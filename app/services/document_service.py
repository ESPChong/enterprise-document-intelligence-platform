"""
Document Processing Service
Handles multi-format document loading, text extraction, and intelligent chunking.
"""

import os
import hashlib
import tempfile
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Any, Optional

from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document

from app.config import settings


class DocumentProcessingError(Exception):
    """Raised when document processing fails."""
    pass


class DocumentService:
    """Service for processing documents into chunks for RAG."""
    
    def __init__(self):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.CHUNK_SIZE,
            chunk_overlap=settings.CHUNK_OVERLAP,
            length_function=len,
            add_start_index=True,
            separators=[
                "\n\n",  # Paragraph breaks
                "\n",    # Line breaks
                ". ",    # Sentence endings
                " ",     # Word boundaries
                ""       # Character fallback
            ]
        )
    
    def validate_file(self, file_path: str) -> bool:
        """Validate if file type is supported and size is acceptable."""
        file_ext = Path(file_path).suffix.lower()
        
        supported_extensions = ['.pdf', '.docx', '.txt']
        if file_ext not in supported_extensions:
            return False
        
        # Check file size
        try:
            file_size_mb = os.path.getsize(file_path) / (1024 * 1024)
            if file_size_mb > settings.MAX_FILE_SIZE_MB:
                return False
        except OSError:
            return False
        
        return True
    
    def generate_document_id(self, file_path: str) -> str:
        """Generate unique ID for document based on content hash."""
        hash_md5 = hashlib.md5()
        with open(file_path, 'rb') as f:
            for chunk in iter(lambda: f.read(4096), b""):
                hash_md5.update(chunk)
        return hash_md5.hexdigest()
    
    def _load_pdf(self, file_path: str) -> List[Document]:
        """Load PDF document using pypdf directly for better control."""
        try:
            from pypdf import PdfReader
            
            reader = PdfReader(file_path)
            documents = []
            
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text.strip():  # Only add non-empty pages
                    documents.append(Document(
                        page_content=text,
                        metadata={
                            "source": file_path,
                            "page": page_num + 1,
                            "total_pages": len(reader.pages)
                        }
                    ))
            
            return documents
            
        except ImportError:
            # Fallback to LangChain loader if pypdf not available
            from langchain_community.document_loaders import PyPDFLoader
            loader = PyPDFLoader(file_path)
            return loader.load()
        except Exception as e:
            raise DocumentProcessingError(f"Failed to load PDF: {str(e)}")
    
    def _load_docx(self, file_path: str) -> List[Document]:
        """Load DOCX document using python-docx directly."""
        try:
            from docx import Document as DocxDocument
            
            doc = DocxDocument(file_path)
            full_text = []
            
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            full_text.append(cell.text)
            
            content = "\n\n".join(full_text)
            
            return [Document(
                page_content=content,
                metadata={"source": file_path}
            )]
            
        except ImportError:
            # Fallback to LangChain loader
            from langchain_community.document_loaders import Docx2txtLoader
            loader = Docx2txtLoader(file_path)
            return loader.load()
        except Exception as e:
            raise DocumentProcessingError(f"Failed to load DOCX: {str(e)}")
    
    def _load_txt(self, file_path: str) -> List[Document]:
        """Load text file with robust encoding handling."""
        # Try multiple encodings
        encodings_to_try = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "ascii"]
        
        for encoding in encodings_to_try:
            try:
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                
                return [Document(
                    page_content=content,
                    metadata={"source": file_path}
                )]
            except (UnicodeDecodeError, UnicodeError):
                continue
        
        # If all encodings fail, read as binary and decode with errors ignored
        try:
            with open(file_path, 'rb') as f:
                content = f.read().decode('utf-8', errors='ignore')
            
            return [Document(
                page_content=content,
                metadata={"source": file_path}
            )]
        except Exception as e:
            raise DocumentProcessingError(f"Failed to load TXT: {str(e)}")
    
    def load_document(self, file_path: str) -> List[Document]:
        """
        Load a document and return list of Document objects with metadata.
        
        Args:
            file_path: Path to the document file
            
        Returns:
            List of Document objects
            
        Raises:
            DocumentProcessingError: If loading fails
        """
        file_ext = Path(file_path).suffix.lower()
        
        if not self.validate_file(file_path):
            raise DocumentProcessingError(
                f"Invalid file: {file_path}. "
                f"Must be PDF, DOCX, or TXT and "
                f"under {settings.MAX_FILE_SIZE_MB}MB"
            )
        
        try:
            # Load based on file type
            if file_ext == ".pdf":
                documents = self._load_pdf(file_path)
            elif file_ext == ".docx":
                documents = self._load_docx(file_path)
            elif file_ext == ".txt":
                documents = self._load_txt(file_path)
            else:
                raise DocumentProcessingError(
                    f"Unsupported file type: {file_ext}"
                )
            
            if not documents:
                raise DocumentProcessingError(
                    f"No content extracted from: {file_path}"
                )
            
            # Add metadata
            doc_id = self.generate_document_id(file_path)
            filename = os.path.basename(file_path)
            
            for doc in documents:
                doc.metadata.update({
                    "document_id": doc_id,
                    "source_file": filename,
                    "file_type": file_ext,
                    "file_size_bytes": os.path.getsize(file_path),
                    "processed_at": datetime.now().isoformat(),
                    "total_pages": len(documents) if file_ext == ".pdf" else 1
                })
            
            return documents
            
        except DocumentProcessingError:
            raise
        except Exception as e:
            raise DocumentProcessingError(f"Failed to load document: {str(e)}")
    
    def chunk_documents(
        self, 
        documents: List[Document],
        custom_chunk_size: Optional[int] = None,
        custom_overlap: Optional[int] = None
    ) -> List[Document]:
        """
        Split documents into overlapping chunks with metadata.
        
        Args:
            documents: List of Document objects
            custom_chunk_size: Override default chunk size
            custom_overlap: Override default overlap
            
        Returns:
            List of chunked Document objects
        """
        # Use custom parameters if provided
        if custom_chunk_size or custom_overlap:
            splitter = RecursiveCharacterTextSplitter(
                chunk_size=custom_chunk_size or settings.CHUNK_SIZE,
                chunk_overlap=custom_overlap or settings.CHUNK_OVERLAP,
                length_function=len,
                add_start_index=True,
                separators=["\n\n", "\n", ". ", " ", ""]
            )
        else:
            splitter = self.text_splitter
        
        all_chunks = []
        
        for doc in documents:
            # Split document into chunks
            chunks = splitter.split_documents([doc])
            
            # Add chunk-specific metadata
            for i, chunk in enumerate(chunks):
                chunk.metadata.update({
                    "chunk_id": f"{doc.metadata['document_id']}_{i}",
                    "chunk_index": i,
                    "total_chunks": len(chunks),
                    "chunk_size": len(chunk.page_content),
                    "start_index": chunk.metadata.get("start_index", 0)
                })
            
            all_chunks.extend(chunks)
        
        return all_chunks
    
    def process_file(self, file_path: str) -> List[Document]:
        """
        Complete pipeline: Load and chunk a document.
        
        Args:
            file_path: Path to document
            
        Returns:
            List of chunked documents ready for embedding
        """
        # Validate file
        if not self.validate_file(file_path):
            raise DocumentProcessingError(
                f"Invalid file: {file_path}. "
                f"Must be PDF, DOCX, or TXT and "
                f"under {settings.MAX_FILE_SIZE_MB}MB"
            )
        
        # Load document
        documents = self.load_document(file_path)
        
        # Chunk documents
        chunks = self.chunk_documents(documents)
        
        return chunks
    
    def get_document_stats(self, chunks: List[Document]) -> Dict[str, Any]:
        """Get statistics about processed chunks."""
        if not chunks:
            return {
                "total_chunks": 0, 
                "total_size_chars": 0, 
                "avg_chunk_size": 0,
                "document_id": None,
                "source_file": None
            }
        
        total_size = sum(len(chunk.page_content) for chunk in chunks)
        
        return {
            "total_chunks": len(chunks),
            "total_size_chars": total_size,
            "avg_chunk_size": total_size // len(chunks) if chunks else 0,
            "document_id": chunks[0].metadata.get("document_id"),
            "source_file": chunks[0].metadata.get("source_file"),
            "processed_at": chunks[0].metadata.get("processed_at")
        }
    
    def get_supported_formats(self) -> List[str]:
        """Get list of supported file formats."""
        return [".pdf", ".docx", ".txt"]